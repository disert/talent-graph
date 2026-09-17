"""岗位需求录入（对话式描述 -> 结构化硬性+择优条件；部门级联（含省份） / 专业 / Excel 批量导入）。"""
from __future__ import annotations

import logging
import threading
import time
import uuid
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Response, UploadFile
from sqlalchemy import select
from sqlalchemy.orm import Session

from .. import timing
from ..concurrency import ProcessingInterrupted, submit_heavy_task
from ..database import get_db
from ..models import JobRequest
from ..schemas import (JobCreate, JobImportItem, JobImportResult, JobImportSessionOut,
                       JobImportStepIn, JobImportStepOut, JobOut)
from ..services import departments, embedding, llm, matching, metrics

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api/jobs", tags=["jobs"])

# Excel 导入/模板允许的后缀
XLSX_SUFFIX = {".xlsx", ".xlsm"}

# 分批导入会话：上传后暂存解析结果，前端逐行调用 step 以便显示进度条。
# 单实例内部工具，用进程内字典即可；带 TTL 与数量上限避免堆积。
_IMPORT_SESSIONS: dict[str, dict] = {}
_IMPORT_LOCK = threading.Lock()
_IMPORT_TTL = 1800          # 未走完的会话保留 30 分钟
_IMPORT_MAX_SESSIONS = 20


# ---------- 工具函数 ----------

def _match_job_task(job_id: int, source: str = "job_create",
                    enqueued_at: float | None = None) -> None:
    """后台任务：岗位入库后与全部在池简历跑两级匹配，结果落库 match_results。

    岗位结构化只花一次 LLM 调用就返回，匹配（多条 LLM 精排）放后台任务队列不阻塞录入。
    enqueued_at 为接口侧提交时刻，用于把「排队等待」也算进耗时汇总。
    """
    from ..database import SessionLocal

    timer = timing.StageTimer(f"岗位匹配 id={job_id} source={source}")
    wait = timing.queue_wait_seconds(enqueued_at)
    if wait is not None:
        timer.record("排队等待", wait)

    db = SessionLocal()
    failed = ""
    display = ""
    try:
        job = db.get(JobRequest, job_id)
        if job is None:
            timer.note("岗位已删除，跳过")
            return
        # 先取纯值：下面的匹配会 rollback 归还连接，ORM 实例随即过期
        display = job.title
        total, records = matching.match_job(db, job, source=source, timer=timer)
        logger.info("岗位 %s 匹配完成：硬性过滤通过 %d 份，写入 %d 条匹配结果",
                    job_id, total, len(records))
    except ProcessingInterrupted as e:
        # 服务正在关闭：本轮匹配作废，不当失败处理
        db.rollback()
        logger.warning("岗位 %s 匹配被服务关闭中断，未写入结果（重启后可在匹配页点「执行匹配」重跑）：%s",
                       job_id, e)
    except Exception as e:
        failed = str(e)[:200]
        db.rollback()
        logger.exception("岗位 %s 匹配失败: %s", job_id, e)
    finally:
        db.close()
        timer.log(error=failed)
        metrics.save(kind="job_match", ref_id=job_id, name=display,
                     metrics=timer.as_dict(), error=failed)


def _embed_or_none(text: str) -> list[float] | None:
    """生成岗位向量；本地无 embedding 模型/接口不可用时降级为 None（仍可入库、匹配）。"""
    try:
        return embedding.embed_one(text)
    except Exception as e:  # 平台接口不可用 / 未装本地模型：不阻断录入
        logger.warning("岗位向量生成失败，已降级为无向量: %s", e)
        return None


def _normalize_department(department: str, department_path: list[str]) -> tuple[str, list[str]]:
    """统一部门级联：优先用路径数组，其次兼容 "/" 分隔字符串 / 单级名称。"""
    path = [str(s).strip() for s in (department_path or []) if str(s).strip()]
    if not path:
        path = departments.department_to_path(department or "")
    if not path:
        return "", []
    return departments.path_to_department(path), list(path)


def _build_job_payload(*, db: Session, title: str, department: str, department_path: list[str],
                       province: str, majors: list[str], raw_text: str,
                       timer: timing.StageTimer | None = None) -> JobRequest:
    """LLM 结构化需求 -> 组装 JobRequest 对象（显式专业并入硬性条件，保证硬过滤生效）。

    province 留空时由所选部门路径在组织架构中的省份推导（子部门默认继承上级）。
    timer 传入时把「大模型结构化」「embedding」两段耗时计入调用方汇总。
    """
    if not title:
        raise ValueError("缺少岗位名称")
    if not raw_text:
        raise ValueError("缺少需求描述")

    with timing.stage(timer, "大模型结构化"):
        result = llm.structure_job_request(raw_text)
    hard = dict(result.get("hard_conditions") or {})
    soft = dict(result.get("soft_conditions") or {})
    summary = str(result.get("summary") or "")

    entered = [m for m in majors if str(m).strip()]
    existing = [str(m) for m in (hard.get("majors") or []) if str(m).strip()]
    merged = list(dict.fromkeys(existing + entered))
    if merged:
        hard["majors"] = merged

    dept_str, dept_path = _normalize_department(department, department_path)
    final_province = (province or "").strip() or departments.resolve_province(db, dept_path)
    with timing.stage(timer, "embedding"):
        embedding_vec = _embed_or_none(embedding.job_to_text(title, soft, summary))
    return JobRequest(
        title=title,
        department=dept_str,
        department_path=dept_path,
        province=final_province,
        majors=list(dict.fromkeys(entered)),
        raw_text=raw_text,
        hard_conditions=hard,
        soft_conditions=soft,
        embedding=embedding_vec,
    )


def _import_one_row(db: Session, item: dict,
                    timer: timing.StageTimer | None = None) -> tuple[bool, int | None, str]:
    """导入单行：LLM 结构化 + 入库 + 排后台匹配。返回 (是否成功, 岗位 id, 失败原因)。"""
    try:
        job = _build_job_payload(
            db=db, title=item.get("title") or "", department=item.get("department") or "",
            department_path=[], province=item.get("province") or "",
            majors=item.get("majors") or [], raw_text=item.get("raw_text") or "",
            timer=timer,
        )
        with timing.stage(timer, "入库"):
            db.add(job)
            db.commit()
            db.refresh(job)
        submit_heavy_task(_match_job_task, job.id, "job_create",
                          enqueued_at=time.perf_counter())
        return True, job.id, ""
    except ValueError as e:      # 行数据不完整（缺岗位名称/需求描述）
        db.rollback()
        return False, None, str(e)
    except RuntimeError as e:    # LLM 结构化失败（连接/解析）
        db.rollback()
        logger.warning("岗位导入失败 第%s行（AI 结构化不可用）: %s", item.get("row"), e)
        return False, None, "AI 结构化服务暂不可用（大模型平台连接失败），该行未入库"
    except Exception as e:       # 单行失败不影响其他行
        db.rollback()
        logger.warning("岗位导入失败 第%s行: %s", item.get("row"), e)
        return False, None, str(e)[:200]


def _purge_import_sessions() -> None:
    """清理过期会话；超过上限时丢弃最早的一个。"""
    now = time.time()
    with _IMPORT_LOCK:
        for token in [t for t, s in _IMPORT_SESSIONS.items() if now - s["created"] > _IMPORT_TTL]:
            _IMPORT_SESSIONS.pop(token, None)
        while len(_IMPORT_SESSIONS) >= _IMPORT_MAX_SESSIONS:
            oldest = min(_IMPORT_SESSIONS, key=lambda t: _IMPORT_SESSIONS[t]["created"])
            _IMPORT_SESSIONS.pop(oldest, None)


def _precheck_items(rows: list[dict]) -> list[dict]:
    """导入前预校验：标出缺岗位名称/需求描述的行，让异常数据在导入前就能看到。"""
    items = []
    for item in rows:
        error = ""
        if not (item.get("title") or "").strip():
            error = "缺少「岗位名称」，该行无法导入"
        elif not (item.get("raw_text") or "").strip():
            error = "缺少「需求描述」，该行无法导入"
        items.append({**item, "error": error})
    return items


def _split_majors(value: object) -> list[str]:
    """拆分 Excel 单元格里的多专业：支持 顿号/逗号/分号/斜杠/空格 等分隔符。"""
    if value is None:
        return []
    text = str(value)
    for sep in ("、", ",", "，", ";", "；", "/", "|", " ", "\n", "\t"):
        text = text.replace(sep, "|")
    return [t.strip() for t in text.split("|") if t.strip()]


def _read_import_rows(file) -> list[dict]:
    """解析第一个工作表：表头 + 数据行。全空行跳过；仅返回原始单元格值，行内校验留到导入循环。"""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(file, data_only=True, read_only=True)
    except Exception as e:
        raise HTTPException(400, f"无法解析 Excel 文件：{e}") from e

    ws = wb.worksheets[0]
    iter_rows = ws.iter_rows(values_only=True)
    try:
        header = next(iter_rows)
    except StopIteration:
        raise HTTPException(400, "Excel 为空，没有表头")

    def find(*keys: str) -> int | None:
        for i, h in enumerate(header):
            if h is None:
                continue
            s = str(h).replace("*", "").replace(" ", "").strip()
            for k in keys:
                if k in s:
                    return i
        return None

    idx_title = find("岗位名称", "岗位")
    idx_raw = find("需求描述", "描述")
    if idx_title is None or idx_raw is None:
        raise HTTPException(400, "Excel 表头缺少必需列：「岗位名称」「需求描述」")

    idx_dept = find("需求部门", "部门")
    idx_prov = find("省份", "省")
    idx_major = find("专业")

    def val(row: tuple, i: int | None) -> object:
        return row[i] if i is not None and i < len(row) else None

    data: list[dict] = []
    for n, row in enumerate(iter_rows, start=2):
        title = str(val(row, idx_title) or "").strip()
        raw = str(val(row, idx_raw) or "").strip()
        if not title and not raw:      # 全空行视为结束/跳过
            continue
        data.append({
            "row": n,
            "title": title,
            "department": str(val(row, idx_dept) or "").strip() if idx_dept is not None else "",
            "province": str(val(row, idx_prov) or "").strip() if idx_prov is not None else "",
            "majors": _split_majors(val(row, idx_major)) if idx_major is not None else [],
            "raw_text": raw,
        })
    if not data:
        raise HTTPException(400, "Excel 中没有可导入的数据行")
    return data


def _build_template_bytes() -> bytes:
    """生成标准导入模板：第 1 个 Sheet 为空白数据表，第 2 个 Sheet 为填写示例。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    headers = [
        "岗位名称*", "需求部门（级联路径，用 / 分隔）", "省份（可选，默认取需求部门所在省份）",
        "专业（多个用顿号或逗号分隔）", "需求描述*",
    ]
    widths = [26, 34, 30, 40, 60]
    fill = PatternFill("solid", fgColor="DDEBF7")
    font = Font(bold=True)

    def style_sheet(ws) -> None:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(vertical="center")

    wb = Workbook()
    ws = wb.active
    ws.title = "导入数据"
    ws.append(headers)
    style_sheet(ws)

    ws2 = wb.create_sheet("填写示例")
    ws2.append(headers)
    style_sheet(ws2)
    ws2.append([
        "新能源并网高级工程师",
        "集团总部/科技创新与数字化部/新能源处",
        "",
        "电气工程、新能源科学与工程、电力系统及其自动化",
        "我们需要一位电力系统自动化方向的博士，35 岁以下，熟悉 PSCAD/MATLAB 仿真，"
        "研究方向偏新能源并网或储能调度，有电网调度或设计院经验优先，能到深圳全职工作。",
    ])
    ws2.append([
        "储能系统研发负责人",
        "产业与新兴业务公司/储能事业部",
        "江苏",
        "储能科学与技术；电化学",
        "负责大型储能电站的系统集成与研发，博士学历，5 年以上储能行业经验，熟悉电池管理系统（BMS）与并网控制。",
    ])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- 接口 ----------

@router.post("", response_model=JobOut)
def create_job(body: JobCreate, db: Session = Depends(get_db)):
    """提交自然语言岗位需求，LLM 结构化为硬性/择优条件并生成 embedding。

    严格失败策略：大模型平台不可达/结构化失败时不入库，返回明确错误，
    避免保存“空条件”岗位干扰后续匹配。入库后后台触发与在池简历的匹配。

    请求内统计「大模型结构化 / embedding / 入库」耗时；匹配在后台任务里另打一行汇总。
    """
    timer = timing.StageTimer(f"岗位录入 title={body.title}")
    try:
        job = _build_job_payload(
            db=db, title=body.title, department=body.department,
            department_path=body.department_path, province=body.province,
            majors=body.majors, raw_text=body.raw_text,
            timer=timer,
        )
    except ValueError as e:  # 缺岗位名称 / 缺需求描述等
        timer.log(error=str(e))
        raise HTTPException(400, str(e)) from e
    except RuntimeError as e:  # LLM 结构化失败（连接/解析）
        logger.warning("岗位 AI 结构化失败，未入库 title=%s: %s", body.title, e)
        timer.log(error=f"AI 结构化失败: {e}")
        raise HTTPException(
            503, "AI 结构化服务暂不可用（大模型平台连接失败），岗位未入库。请稍后重试，或检查 backend/.env 的大模型配置。"
        ) from e
    except Exception as e:
        logger.exception("岗位结构化异常 title=%s", body.title)
        timer.log(error=str(e))
        raise HTTPException(502, f"岗位结构化失败，岗位未入库：{e}") from e

    with timer.stage("入库"):
        db.add(job)
        db.commit()
        db.refresh(job)
    submit_heavy_task(_match_job_task, job.id, "job_create",
                      enqueued_at=time.perf_counter())  # 后台匹配在池简历
    timer.note(f"id={job.id}")
    timer.log()
    metrics.save(kind="job_create", ref_id=job.id, name=job.title, metrics=timer.as_dict())
    return job


@router.get("/departments")
def list_departments(db: Session = Depends(get_db)):
    """组织架构树（岗位级联选择数据源，读自 departments 表）。"""
    return departments.load_tree(db)


@router.post("/import", response_model=JobImportResult)
def import_jobs(file: UploadFile, db: Session = Depends(get_db)):
    """Excel 批量导入岗位（一次性提交版）。逐行容错：单行失败不影响其他行；LLM 逐条结构化。

    前端「导入 Excel」走 /import/prepare + /import/step 以便显示进度条；
    本接口保留给脚本/接口直调，行为与原实现一致（一次性返回汇总）。
    """
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in XLSX_SUFFIX:
        raise HTTPException(400, "仅支持 .xlsx / .xlsm 格式（可先下载模板）")
    rows = _read_import_rows(file.file)

    batch_timer = timing.StageTimer(f"岗位导入(一次性) file={file.filename} rows={len(rows)}")
    result = JobImportResult(total=len(rows))
    for item in rows:
        item_timer = timing.StageTimer(f"岗位导入 第{item['row']}行 title={item.get('title')}")
        ok, job_id, error = _import_one_row(db, item, timer=item_timer)
        if ok:
            result.succeeded += 1
            item_timer.note(f"id={job_id}")
        else:
            result.failed += 1
            result.errors.append({"row": item["row"], "message": error})
        item_timer.log(error=error)
        metrics.save(kind="job_create", ref_id=job_id or 0, name=item.get("title") or "",
                     metrics=item_timer.as_dict(), error=error)
    batch_timer.note(f"成功{result.succeeded} 失败{result.failed}")
    batch_timer.log()
    return result


@router.post("/import/prepare", response_model=JobImportSessionOut)
def prepare_import(file: UploadFile, db: Session = Depends(get_db)):
    """解析 Excel 并建立导入会话：返回逐行清单（含预校验异常），供前端渲染进度条。

    只解析、不入库；入库由 /import/step 逐行触发，前端据此显示实时进度与异常明细。
    """
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in XLSX_SUFFIX:
        raise HTTPException(400, "仅支持 .xlsx / .xlsm 格式（可先下载模板）")
    rows = _read_import_rows(file.file)
    items = _precheck_items(rows)

    token = uuid.uuid4().hex
    _purge_import_sessions()
    with _IMPORT_LOCK:
        _IMPORT_SESSIONS[token] = {"created": time.time(), "items": items}

    return JobImportSessionOut(
        token=token,
        total=len(items),
        valid=sum(1 for i in items if not i["error"]),
        invalid=sum(1 for i in items if i["error"]),
        items=[JobImportItem(**i) for i in items],
    )


@router.post("/import/step", response_model=JobImportStepOut)
def import_step(body: JobImportStepIn, db: Session = Depends(get_db)):
    """导入会话中的第 index 行：入库并排后台匹配，返回该行结果。

    前端循环调用以推进进度条；每行独立事务，失败不阻塞后续行。
    """
    with _IMPORT_LOCK:
        session = _IMPORT_SESSIONS.get(body.token)
        items = list(session["items"]) if session else None
    if session is None or items is None:
        raise HTTPException(400, "导入会话已过期，请重新选择文件")
    if not 0 <= body.index < len(items):
        raise HTTPException(400, f"导入行序号越界：{body.index}")

    item = items[body.index]
    # 每行一次日志：前端进度条 + 后端逐行耗时，批量导入慢时能一眼看出卡在哪一行
    item_timer = timing.StageTimer(f"岗位导入 第{item['row']}行 title={item.get('title')}")
    ok, job_id, error = _import_one_row(db, item, timer=item_timer)
    if ok:
        item_timer.note(f"id={job_id}")
    item_timer.log(error=error)
    metrics.save(kind="job_create", ref_id=job_id or 0, name=item.get("title") or "",
                 metrics=item_timer.as_dict(), error=error)
    return JobImportStepOut(row=item["row"], title=item.get("title") or "",
                            ok=ok, job_id=job_id, error=error)


@router.get("/import/template")
def download_import_template():
    """下载岗位批量导入的标准 Excel 模板。"""
    filename = quote("岗位需求导入模板.xlsx")
    return Response(
        content=_build_template_bytes(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


@router.get("", response_model=list[JobOut])
def list_jobs(db: Session = Depends(get_db)):
    return list(db.scalars(select(JobRequest).order_by(JobRequest.created_at.desc())))


@router.get("/{job_id}", response_model=JobOut)
def get_job(job_id: int, db: Session = Depends(get_db)):
    job = db.get(JobRequest, job_id)
    if job is None:
        raise HTTPException(404, "岗位不存在")
    return job


@router.delete("/{job_id}")
def delete_job(job_id: int, db: Session = Depends(get_db)):
    job = db.get(JobRequest, job_id)
    if job is None:
        raise HTTPException(404, "岗位不存在")
    db.delete(job)
    db.commit()
    return {"ok": True}
